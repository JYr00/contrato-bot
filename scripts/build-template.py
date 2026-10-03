"""
Genera templates/contrato-arrendamiento.docx a partir del texto del contrato original.

Los campos variables usan la sintaxis de docxtemplater: {nombre_campo}.
Cada placeholder queda dentro de un solo "run" para que docxtemplater lo reconozca.

Uso:  python3 scripts/build-template.py   (requiere: pip install python-docx)
Tras generarla, la plantilla se puede editar a mano en Word sin problema,
siempre que no se partan los {placeholders}.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent.parent / "templates" / "contrato-arrendamiento.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # Carta
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))

def campo(par, codigo):
    """Inserta un campo de Word (PAGE, NUMPAGES) que el procesador calcula al abrir o convertir."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), codigo)
    run = OxmlElement("w:r")
    texto = OxmlElement("w:t")
    texto.text = "1"
    run.append(texto)
    fld.append(run)
    par._p.append(fld)


pie = sec.footer.paragraphs[0]
pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
pie.add_run("Página ")
campo(pie, "PAGE")
pie.add_run(" de ")
campo(pie, "NUMPAGES")
for run in pie.runs:
    run.font.size = Pt(8)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15


def p(title: str = "", body: str = "", *, indent: float = 0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    par = doc.add_paragraph()
    par.alignment = align
    if indent:
        par.paragraph_format.left_indent = Cm(indent)
    if title:
        par.add_run(title).bold = True
    if body:
        par.add_run(body)
    return par


def mantener_junta(tabla):
    """Evita que la tabla se parta entre páginas: ninguna fila se divide y cada fila va con la siguiente."""
    for i, fila in enumerate(tabla.rows):
        trPr = fila._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for celda in fila.cells:
            for par in celda.paragraphs:
                par.paragraph_format.keep_together = True
                par.paragraph_format.keep_with_next = i < len(tabla.rows) - 1


# --- Encabezado -------------------------------------------------------------
t = p("CONTRATO DE ARRENDAMIENTO DE VIVIENDA URBANA", align=WD_ALIGN_PARAGRAPH.CENTER)
t.runs[0].font.size = Pt(12)
t.paragraph_format.space_after = Pt(12)

p("", "MARIO A. ROJAS RODELO, con domicilio en la ciudad de BOGOTÁ, identificado con cédula de "
  "ciudadanía No. 80.101.225, quien obra en nombre propio y que para efectos de este contrato se "
  "denominará “EL ARRENDADOR”, por una parte, y por la otra, {arrendatarios_texto} EL ARRENDATARIO, "
  "manifestaron que han decidido celebrar un contrato de arrendamiento "
  "de bien inmueble destinado a vivienda, en adelante el “Contrato”, el cual se rige por la Ley 820 "
  "de 2003 y por las siguientes cláusulas:")

p("Primera. – Objeto: ",
  "Por medio del presente contrato, EL ARRENDADOR entrega a título de arrendamiento a EL "
  "ARRENDATARIO el siguiente bien inmueble: {inmueble_direccion}, destinado "
  "para el uso de vivienda para {ocupantes_texto}. LINDEROS: Se plasmarán en documento anexo que "
  "hará parte del contrato.")

p("Segunda. – Canon de Arrendamiento: ",
  "El canon de arrendamiento mensual es la suma de {precio_texto} M/cte, que EL ARRENDATARIO pagará "
  "anticipadamente al ARRENDADOR o a su orden, en el domicilio de EL ARRENDADOR ubicado en la "
  "Carrera 105 H 67 D 33, dentro de los primeros cinco (5) días de cada mes. Cada doce (12) meses el "
  "canon de arrendamiento será reajustado en el porcentaje máximo correspondiente al IPC del año "
  "inmediatamente anterior.")
p("Parágrafo 1: ",
  "La tolerancia de EL ARRENDADOR en recibir el pago del canon de arrendamiento con posterioridad "
  "al plazo indicado para ello en esta Cláusula no podrá entenderse, en ningún caso, como ánimo de "
  "EL ARRENDADOR de modificar el término establecido en este Contrato para el pago del canon.")
# Párrafo condicional: docxtemplater (paragraphLoop) elimina los párrafos de apertura y cierre,
# y omite el parágrafo completo cuando no se pactó depósito.
p("", "{#hay_deposito}")
p("Parágrafo 2: ",
  "A la firma del presente Contrato, EL ARRENDATARIO entrega a EL ARRENDADOR la suma de "
  "{deposito_texto} M/cte, a título de depósito, para cubrir los daños al inmueble, los servicios "
  "públicos y demás obligaciones pendientes a cargo de EL ARRENDATARIO a la fecha de restitución del "
  "Inmueble. El saldo, si lo hubiere, será devuelto a EL ARRENDATARIO una vez restituido el Inmueble "
  "y verificado el pago de dichas obligaciones.")
p("", "{/hay_deposito}")

p("Tercera. – Vigencia: ",
  "El arrendamiento tendrá una duración de {duracion_texto}, contados a partir del "
  "{fecha_inicio_texto}. No obstante lo anterior, el término de este contrato no se prorrogará "
  "automáticamente.")

p("Cuarta. – Entrega: ",
  "EL ARRENDATARIO, en la fecha de suscripción de este documento, declara recibir el inmueble de "
  "manos de EL ARRENDADOR en perfecto estado, de conformidad con el inventario elaborado por las "
  "partes y que forma parte integrante de este contrato.")

p("Quinta. – Reparación: ",
  "Los daños que se ocasionen al inmueble por EL ARRENDATARIO, por responsabilidad suya o de sus "
  "dependientes, serán reparados y cubiertos sus costos de reparación en su totalidad por EL "
  "ARRENDATARIO. Igualmente, EL ARRENDATARIO se obliga a cumplir con las obligaciones previstas en "
  "los artículos 2029 y 2030 del Código Civil.")
p("Parágrafo: ",
  "EL ARRENDATARIO se abstendrá de hacer mejoras de cualquier clase al inmueble sin permiso previo "
  "y escrito de EL ARRENDADOR. Las mejoras al inmueble serán del propietario del inmueble y no habrá "
  "lugar al reconocimiento del precio, costo o indemnización alguna al ARRENDATARIO por las mejoras "
  "realizadas. Las mejoras no podrán retirarse salvo que EL ARRENDADOR lo exija por escrito, a lo que "
  "EL ARRENDATARIO accederá inmediatamente a su costa, dejando el inmueble en el mismo buen estado en "
  "que lo recibió de EL ARRENDADOR, salvo el deterioro natural por el uso legítimo.")

p("Sexta. – Servicios Públicos: ",
  "EL ARRENDATARIO pagará oportuna y totalmente los servicios públicos del inmueble: LUZ, GAS, "
  "ACUEDUCTO, ALCANTARILLADO Y ASEO.")
p("Parágrafo 1: ",
  "EL ARRENDATARIO declara que ha recibido en perfecto estado de funcionamiento y de conservación "
  "las instalaciones para uso de los servicios públicos del Inmueble, que se abstendrá de "
  "modificarlas sin permiso previo y escrito de EL ARRENDADOR y que responderá por daños y/o "
  "violaciones de los reglamentos de las correspondientes empresas de servicios públicos.")
p("Parágrafo 2: ",
  "EL ARRENDATARIO reconoce que EL ARRENDADOR en ningún caso y en ninguna circunstancia es "
  "responsable por la interrupción o deficiencia en la prestación de cualquiera de los servicios "
  "públicos del Inmueble. En caso de la prestación deficiente o suspensión de cualquiera de los "
  "servicios públicos del Inmueble, EL ARRENDATARIO reclamará de manera directa a las empresas "
  "prestadoras del servicio y no a EL ARRENDADOR.")

p("Séptima. – Destinación: ",
  "EL ARRENDATARIO, durante la vigencia del Contrato, destinará el Inmueble única y exclusivamente "
  "para su vivienda y la de su familia. En ningún caso EL ARRENDATARIO podrá subarrendar o ceder en "
  "todo o en parte este arrendamiento, so pena de que EL ARRENDADOR pueda dar por terminado "
  "válidamente el Contrato en forma inmediata, sin lugar a indemnización alguna en favor de EL "
  "ARRENDATARIO, y podrá exigir la devolución del Inmueble sin necesidad de ningún tipo de "
  "requerimiento previo por parte de EL ARRENDADOR. Igualmente, EL ARRENDATARIO se abstendrá de "
  "guardar o permitir que dentro del Inmueble se guarden semovientes o animales domésticos y/o "
  "elementos inflamables, tóxicos, insalubres, explosivos o dañosos para la conservación, higiene, "
  "seguridad y estética del inmueble y en general de sus ocupantes permanentes o transitorios.")
p("Parágrafo: ",
  "EL ARRENDADOR declara expresa y terminantemente prohibida la destinación del inmueble a los fines "
  "contemplados en el literal b) del parágrafo del Artículo 34 de la Ley 30 de 1986 y, en "
  "consecuencia, EL ARRENDATARIO se obliga a no usar el Inmueble para el ocultamiento de personas, "
  "depósito de armas o explosivos y dinero de los grupos terroristas. No destinará el inmueble para "
  "la elaboración, almacenamiento o venta de sustancias alucinógenas tales como marihuana, hachís, "
  "cocaína, metacualona y similares. EL ARRENDATARIO faculta a EL ARRENDADOR para que, directamente "
  "o a través de sus funcionarios debidamente autorizados por escrito, visite el Inmueble para "
  "verificar el cumplimiento de las obligaciones de EL ARRENDATARIO.")

p("Octava. – Restitución: ",
  "Terminado el contrato en los términos establecidos en el presente documento y de conformidad con "
  "la ley, EL ARRENDATARIO restituirá el Inmueble a EL ARRENDADOR en las mismas buenas condiciones en "
  "que lo recibió de EL ARRENDADOR, salvo el deterioro natural causado por el uso legítimo, y "
  "entregará a EL ARRENDADOR los ejemplares originales de las facturas de cobro por concepto de "
  "servicios públicos del Inmueble correspondientes a los últimos tres (3) meses, debidamente "
  "canceladas por EL ARRENDATARIO, bajo el entendido que hará entrega de dichas facturas en el "
  "domicilio de EL ARRENDADOR con una antelación de cinco (5) días hábiles a la fecha fijada para la "
  "restitución material del Inmueble al ARRENDADOR.")
p("Parágrafo 1: ",
  "No obstante lo anterior, EL ARRENDADOR podrá negarse a recibir el Inmueble cuando a su juicio "
  "existan obligaciones pendientes a cargo de EL ARRENDATARIO que no hayan sido satisfechas en forma "
  "debida, caso en el cual se seguirá causando el canon de arrendamiento hasta que EL ARRENDATARIO "
  "cumpla con lo que le corresponde.")
p("Parágrafo 2: ",
  "La responsabilidad de EL ARRENDATARIO subsistirá aún después de restituido el Inmueble, mientras "
  "EL ARRENDADOR no haya entregado el paz y salvo correspondiente por escrito a EL ARRENDATARIO.")
p("Parágrafo 3: ",
  "Independientemente de la vigencia del contrato, EL ARRENDATARIO podrá hacer entrega del inmueble "
  "en cualquier momento avisando a EL ARRENDADOR con ocho (8) días de anticipación.")

p("Novena. – Incumplimiento: ",
  "El incumplimiento por parte de EL ARRENDATARIO de cualquiera de las obligaciones a su cargo "
  "contenidas en la ley o en este Contrato faculta a EL ARRENDADOR para ejercer las siguientes "
  "acciones, simultáneamente o en el orden que este elija:")
p("", "a) Declarar terminado este Contrato y reclamar la devolución del Inmueble judicial y/o "
  "extrajudicialmente;", indent=1)
p("", "b) Exigir y perseguir a través de cualquier medio, judicial o extrajudicialmente, de EL "
  "ARRENDATARIO el monto de los perjuicios resultantes del incumplimiento, así como de la multa por "
  "incumplimiento pactada en este Contrato.", indent=1)
p("Parágrafo: Terminación del Contrato de Arrendamiento.")
p("Terminación por mutuo acuerdo. ",
  "Las partes, en cualquier tiempo y de común acuerdo, podrán dar por terminado el contrato de "
  "vivienda urbana.")
p("Terminación por parte del arrendador. ",
  "Son causales para que el arrendador pueda pedir unilateralmente la terminación del contrato las "
  "siguientes:")
for item in [
    "1. La no cancelación por parte del arrendatario de las rentas y reajustes dentro del término "
    "estipulado en el contrato.",
    "2. La no cancelación de los servicios públicos que cause la desconexión o pérdida del servicio, "
    "o el pago de las expensas comunes cuando su pago estuviere a cargo del arrendatario.",
    "3. El subarriendo total o parcial del inmueble, la cesión del contrato o del goce del inmueble o "
    "el cambio de destinación del mismo por parte del arrendatario, sin expresa autorización del "
    "arrendador.",
    "4. La incursión reiterada del arrendatario en procederes que afecten la tranquilidad ciudadana "
    "de los vecinos, o la destinación del inmueble para actos delictivos o que impliquen "
    "contravención, debidamente comprobados ante la autoridad policiva.",
    "5. La realización de mejoras, cambios o ampliaciones del inmueble sin expresa autorización del "
    "arrendador, o la destrucción total o parcial del inmueble o área arrendada por parte del "
    "arrendatario.",
    "6. El arrendador podrá dar por terminado unilateralmente el contrato de arrendamiento a la fecha "
    "de vencimiento del término inicial o de sus prórrogas invocando cualquiera de las siguientes "
    "causales especiales de restitución, previo aviso escrito al arrendatario a través del servicio "
    "postal autorizado, con una antelación no menor a tres (3) meses a la referida fecha de "
    "vencimiento:",
]:
    p("", item, indent=0.75)
for item in [
    "a) Cuando el propietario o poseedor del inmueble necesitare ocuparlo para su propia habitación, "
    "por un término no menor de un (1) año;",
    "b) Cuando el inmueble haya de demolerse para efectuar una nueva construcción, o cuando se "
    "requiera desocuparlo con el fin de ejecutar obras independientes para su reparación;",
    "c) Cuando haya de entregarse en cumplimiento de las obligaciones originadas en un contrato de "
    "compraventa;",
    "d) La plena voluntad de dar por terminado el contrato, siempre y cuando el contrato de "
    "arrendamiento cumpliere como mínimo cuatro (4) años de ejecución. El arrendador deberá "
    "indemnizar al arrendatario con una suma equivalente al precio de tres (3) meses de "
    "arrendamiento.",
]:
    p("", item, indent=1.5)
p("Terminación por parte del arrendatario. ",
  "Son causales para que el arrendatario pueda pedir unilateralmente la terminación del contrato las "
  "siguientes:")
for item in [
    "1. La suspensión de la prestación de los servicios públicos al inmueble, por acción premeditada "
    "del arrendador o porque incurra en mora en pagos que estuvieren a su cargo.",
    "2. La incursión reiterada del arrendador en procederes que afecten gravemente el disfrute cabal "
    "por el arrendatario del inmueble arrendado, debidamente comprobada ante la autoridad policiva.",
    "3. El desconocimiento por parte del arrendador de derechos reconocidos al arrendatario por la ley "
    "o contractualmente.",
]:
    p("", item, indent=0.75)

p("Décima. – Cesión: ",
  "EL ARRENDADOR podrá ceder libremente, total o parcialmente, los derechos que se derivan de este "
  "Contrato, y dicha cesión producirá efectos respecto de EL ARRENDATARIO a partir de la fecha de la "
  "comunicación escrita en la que se le notifique la cesión. EL ARRENDATARIO no podrá ceder este "
  "Contrato sin autorización previa y escrita de EL ARRENDADOR.")

p("Décima Primera. – Renuncia a requerimientos: ",
  "EL ARRENDATARIO renuncia expresamente a los requerimientos privados y judiciales previstos en la "
  "ley para efectos de su constitución en mora en el pago del canon de arrendamiento o de cualquier "
  "otra obligación derivada de este Contrato.")

p("Décima Segunda. – Validez: ",
  "El presente Contrato anula todo convenio anterior relativo al arrendamiento del mismo Inmueble y "
  "solamente podrá ser modificado por escrito suscrito por las Partes.")

p("Décima Tercera. – Mérito Ejecutivo: ",
  "EL ARRENDATARIO declara de manera expresa que reconoce y acepta que este Contrato presta mérito "
  "ejecutivo para exigir de EL ARRENDATARIO y a favor de EL ARRENDADOR el pago de los cánones de "
  "arrendamiento causados y no pagados por EL ARRENDATARIO, las multas y sanciones que se causen por "
  "el incumplimiento de EL ARRENDATARIO de cualquiera de las obligaciones a su cargo en virtud de la "
  "ley o de este Contrato, las sumas causadas y no pagadas por EL ARRENDATARIO por concepto de "
  "servicios públicos del Inmueble, y cualquier otra suma de dinero que por cualquier concepto deba "
  "ser pagada por EL ARRENDATARIO, para lo cual bastará la sola afirmación de incumplimiento de EL "
  "ARRENDATARIO hecha por EL ARRENDADOR, afirmación que solo podrá ser desvirtuada por EL "
  "ARRENDATARIO con la presentación de los respectivos recibos de pago.")
p("Parágrafo: ",
  "Las Partes acuerdan que cualquier copia autenticada ante Notario de este Contrato tendrá el mismo "
  "valor que el original para efectos judiciales y extrajudiciales.")

p("Décima Cuarta. – Costos: ",
  "Cualquier costo que se cause con ocasión de la celebración o prórroga de este Contrato, incluyendo "
  "el impuesto de timbre, será asumido en su integridad por EL ARRENDATARIO.")

p("Décima Quinta. – Cláusula Penal: ",
  "En el evento de incumplimiento de cualquiera de las Partes a las obligaciones a su cargo "
  "contenidas en la ley o en este Contrato, la parte incumplida deberá pagar a la otra parte una suma "
  "equivalente a dos (2) cánones de arrendamiento vigentes en la fecha del incumplimiento, a título "
  "de pena. En el evento que los perjuicios ocasionados por la parte incumplida excedan el valor de "
  "la suma aquí prevista como pena, la parte incumplida deberá pagar a la otra parte la diferencia "
  "entre el valor total de los perjuicios y el valor de la pena prevista en esta Cláusula.")

p("Décima Sexta. – Autorización: ",
  "EL ARRENDATARIO autoriza expresa e irrevocablemente a EL ARRENDADOR y/o al cesionario de este "
  "Contrato a consultar información de EL ARRENDATARIO que obre en las bases de datos de información "
  "del comportamiento financiero y crediticio o centrales de riesgo que existan en el país, así como "
  "a reportar a dichas bases de datos cualquier incumplimiento de EL ARRENDATARIO a este Contrato.")

p("Décima Séptima. – Abandono: ",
  "EL ARRENDATARIO autoriza de manera expresa e irrevocable a EL ARRENDADOR para ingresar al Inmueble "
  "y recuperar su tenencia, con el solo requisito de la presencia de dos (2) testigos, en procura de "
  "evitar el deterioro o desmantelamiento del Inmueble, en el evento que por cualquier causa o "
  "circunstancia el Inmueble permanezca abandonado o deshabitado por el término de un (1) mes o más y "
  "que la exposición al riesgo sea tal que amenace la integridad física del bien o la seguridad del "
  "vecindario.")

p("Décima Octava. – Recibos de pago de servicios públicos: ",
  "EL ARRENDADOR, en cualquier tiempo durante la vigencia de este Contrato, podrá exigir a EL "
  "ARRENDATARIO la presentación de las facturas de los servicios públicos del Inmueble a fin de "
  "verificar la cancelación de los mismos. En el evento que EL ARRENDADOR llegare a comprobar que "
  "alguna de las facturas no ha sido pagada por EL ARRENDATARIO encontrándose vencido el plazo para "
  "su pago, EL ARRENDADOR podrá pagarla directamente y EL ARRENDATARIO deberá reembolsarle su valor "
  "dentro de los cinco (5) días siguientes al requerimiento que para tal efecto se le haga, sin "
  "perjuicio de que dicho incumplimiento constituya causal de terminación del Contrato.")

# Desde la cláusula de notificaciones hasta las firmas va todo encadenado (keep-with-next): si no cabe, pasa
# completo a la página siguiente. Así las firmas nunca quedan solas en una hoja sin texto del contrato.
notificaciones = p("Décima Novena. – Notificaciones: ",
  "Para todos los efectos de este Contrato, las Partes recibirán notificaciones en las siguientes "
  "direcciones:")
notificaciones.paragraph_format.keep_with_next = True

tabla = doc.add_table(rows=4, cols=2)
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
filas = [
    ("ARRENDATARIO", "ARRENDADOR"),
    ("{arrendatario_nombre}\n{arrendatario_direccion}", "MARIO A. ROJAS RODELO\nCarrera 105 H 67 D 33"),
    ("Correo electrónico: {arrendatario_correo}", "Correo electrónico: {arrendador_correo}"),
    ("Celular: {arrendatario_celular}", "Celular: {arrendador_celular}"),
]
for r, (izq, der) in enumerate(filas):
    for c, texto in enumerate((izq, der)):
        cell = tabla.cell(r, c)
        cell.paragraphs[0].text = ""
        lineas = texto.split("\n")
        for i, linea in enumerate(lineas):
            par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            par.paragraph_format.space_after = Pt(2)
            run = par.add_run(linea)
            run.bold = r == 0
mantener_junta(tabla)
for celda in tabla.rows[-1].cells:  # la última fila también sigue con el cierre
    for par in celda.paragraphs:
        par.paragraph_format.keep_with_next = True

doc.add_paragraph().paragraph_format.keep_with_next = True
constancia = p("", "Para constancia, el presente Contrato es suscrito en la ciudad de Bogotá, en "
  "{ejemplares_texto} ejemplares de igual valor, cada uno de ellos con destino a cada una de las "
  "partes.")
constancia.paragraph_format.keep_with_next = True
constancia.paragraph_format.keep_together = True

for _ in range(2):
    doc.add_paragraph().paragraph_format.keep_with_next = True
firmas = doc.add_table(rows=4, cols=2)
firmas.alignment = WD_TABLE_ALIGNMENT.CENTER
contenido = [
    ("______________________________", "______________________________"),
    ("EL ARRENDADOR", "EL ARRENDATARIO"),
    ("MARIO A. ROJAS RODELO", "{arrendatario_nombre}"),
    ("C.C. 80.101.225", "{arrendatario_documento_firma}"),
]
for r, (izq, der) in enumerate(contenido):
    for c, texto in enumerate((izq, der)):
        par = firmas.cell(r, c).paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        run = par.add_run(texto)
        run.bold = r == 1
mantener_junta(firmas)

# Firmas de los co-arrendatarios: el bloque se repite por cada uno (docxtemplater elimina los párrafos
# de apertura y cierre del ciclo) y queda junto a la tabla de firmas.
p("", "{#coarrendatarios}").paragraph_format.keep_with_next = True
for texto, negrita in [
    ("", False),
    ("", False),
    ("______________________________", False),
    ("EL ARRENDATARIO", True),
    ("{nombre}", False),
    ("{documento_firma}", False),
]:
    par = p("", align=WD_ALIGN_PARAGRAPH.CENTER)
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.keep_with_next = True
    if texto:
        par.add_run(texto).bold = negrita
p("", "{/coarrendatarios}")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Plantilla generada en {OUT}")
